"""Excel export of a project's full document set.

Its own concern alongside db.py (storage) and extract.py (extraction) — main.py
just calls build_workbook() and streams the result.
"""

from __future__ import annotations

import io
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

NUMERIC_FORMAT = "#,##0.00"

DOC_COLUMNS = [
    ("document_id", "Document ID"), ("document_type", "Type"), ("status", "Status"),
    ("doc_number", "Doc Number"), ("po_number", "PO Number"), ("dc_number", "DC Number"),
    ("doc_date", "Doc Date"), ("vendor_name_raw", "Vendor"), ("vendor_gstin", "Vendor GSTIN"),
    ("place_of_supply", "Place of Supply"), ("basic_value", "Basic Value"),
    ("tax_type", "Tax Type"), ("igst_amount", "IGST"), ("cgst_amount", "CGST"),
    ("sgst_amount", "SGST"), ("tcs_amount", "TCS"), ("rounding_off", "Rounding Off"),
    ("total_value", "Total Value"), ("uploaded_at", "Uploaded At"),
    ("reviewed_by", "Reviewed By"), ("reviewed_at", "Reviewed At"),
]
DOC_NUMERIC_KEYS = {
    "basic_value", "igst_amount", "cgst_amount", "sgst_amount",
    "tcs_amount", "rounding_off", "total_value",
}

LINE_COLUMNS = [
    ("document_id", "Document ID"), ("doc_number", "Doc Number"), ("line_no", "Line No"),
    ("description_raw", "Description"), ("material", "Material"), ("hsn_code", "HSN Code"),
    ("quantity", "Quantity"), ("unit", "Unit"), ("rate", "Rate"), ("amount", "Amount"),
    ("tax_rate", "Tax Rate %"), ("dc_number", "DC Number"), ("dc_date", "DC Date"),
]
LINE_NUMERIC_KEYS = {"quantity", "rate", "amount", "tax_rate"}


def build_workbook(con: sqlite3.Connection, project: sqlite3.Row) -> io.BytesIO:
    """4 sheets: Summary, Documents, Line Items, Materials Rollup.

    Rejected documents are excluded from all of them. This file is a statement
    of what the project bought, and it leaves the building — a refused invoice
    sitting in it, even flagged, is one unread Status column away from being
    read as a real purchase. The count of what was dropped goes on Summary so
    the omission is visible rather than silent; the documents themselves stay
    in the console, which is where the audit trail lives.

    A scanned batch the console hasn't yet chosen Process or Draft for
    (awaiting_scan_decision) is excluded outright, same as list_documents —
    it isn't part of the project's documents yet, so it isn't part of what
    left the building either.
    """
    every_document = con.execute(
        "SELECT * FROM documents WHERE project_id = ? AND awaiting_scan_decision = 0"
        " ORDER BY uploaded_at DESC",
        (project["id"],),
    ).fetchall()
    rejected_count = sum(1 for d in every_document if d["status"] == "REJECTED")
    documents = [d for d in every_document if d["status"] != "REJECTED"]
    doc_ids = [d["id"] for d in documents]

    headers, lines = {}, []
    if doc_ids:
        placeholders = ",".join("?" * len(doc_ids))
        headers = {
            h["document_id"]: h
            for h in con.execute(
                f"SELECT * FROM doc_headers WHERE document_id IN ({placeholders})", doc_ids
            ).fetchall()
        }
        lines = con.execute(
            f"SELECT * FROM doc_lines WHERE document_id IN ({placeholders}) "
            "ORDER BY document_id, line_no",
            doc_ids,
        ).fetchall()

    materials = {m["id"]: m["name"] for m in con.execute("SELECT id, name FROM materials")}

    def material_label(line: sqlite3.Row) -> str:
        return materials.get(line["material_id"]) or line["description_raw"] or "Unclassified"

    # `documents`, `headers` and `lines` are already rejection-free — the
    # filter above is the only place that rule is applied.
    wb = Workbook()
    _write_summary(wb.active, project, documents, headers, rejected_count)
    _write_documents(wb.create_sheet("Documents"), documents, headers)
    _write_lines(wb.create_sheet("Line Items"), lines, headers, material_label)
    _write_rollup(wb.create_sheet("Materials Rollup"), lines, material_label)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _header_row(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _format_numeric(ws, indexes: set[int]) -> None:
    for row in ws.iter_rows(min_row=2):
        for idx in indexes:
            row[idx - 1].number_format = NUMERIC_FORMAT


def _autosize(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 40)


def _write_summary(ws, project, documents, headers, rejected_count=0) -> None:
    ws.title = "Summary"
    by_type, by_status = {}, {}
    for d in documents:
        by_type[d["document_type"]] = by_type.get(d["document_type"], 0) + 1
        by_status[d["status"]] = by_status.get(d["status"], 0) + 1
    # `documents` and `headers` arrive already free of rejections.
    total_value = sum(h["total_value"] or 0 for h in headers.values())

    rows = [
        ("Project Code", project["code"]), ("Project Name", project["name"]),
        ("Client", project["client"]), ("Location", project["location"]),
        ("Site", project["site"]), ("Status", project["status"]),
        ("Created At", project["created_at"]), ("", ""),
        ("Total Documents", len(documents)),
    ]
    rows += [(f"— {k}", v) for k, v in sorted(by_type.items())]
    rows += [(f"— {k}", v) for k, v in sorted(by_status.items())]
    rows.append(("Total Value", round(total_value, 2)))
    # Stated, not hidden: without this line the document count here silently
    # disagrees with the console and nothing explains why.
    if rejected_count:
        rows += [("", ""), ("Rejected (excluded)", rejected_count)]

    for label, value in rows:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28


def _write_documents(ws, documents, headers) -> None:
    _header_row(ws, [label for _, label in DOC_COLUMNS])
    for d in documents:
        h = headers.get(d["id"])
        row = []
        for key, _ in DOC_COLUMNS:
            if key == "document_id":
                row.append(d["id"])
            elif key in ("document_type", "status", "uploaded_at"):
                row.append(d[key])
            else:
                row.append(h[key] if h else None)
        ws.append(row)
    _format_numeric(ws, {i + 1 for i, (k, _) in enumerate(DOC_COLUMNS) if k in DOC_NUMERIC_KEYS})
    _autosize(ws)


def _write_lines(ws, lines, headers, material_label) -> None:
    _header_row(ws, [label for _, label in LINE_COLUMNS])
    for line in lines:
        h = headers.get(line["document_id"])
        ws.append([
            line["document_id"], h["doc_number"] if h else None, line["line_no"],
            line["description_raw"], material_label(line), line["hsn_code"],
            line["quantity"], line["unit"], line["rate"], line["amount"],
            line["tax_rate"], line["dc_number"], line["dc_date"],
        ])
    _format_numeric(ws, {i + 1 for i, (k, _) in enumerate(LINE_COLUMNS) if k in LINE_NUMERIC_KEYS})
    _autosize(ws)


def _write_rollup(ws, lines, material_label) -> None:
    _header_row(ws, ["Material", "Unit", "Quantity", "Docs"])
    rollup = {}
    for line in lines:
        key = (material_label(line), line["unit"] or "")
        row = rollup.setdefault(key, {"qty": 0.0, "docs": set()})
        row["qty"] += line["quantity"] or 0
        row["docs"].add(line["document_id"])
    for (material, unit), row in sorted(rollup.items(), key=lambda kv: -kv[1]["qty"]):
        ws.append([material, unit or "—", round(row["qty"], 2), len(row["docs"])])
    _format_numeric(ws, {3})
    _autosize(ws)


if __name__ == "__main__":
    import db as _db
    from openpyxl import load_workbook

    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript(_db.SCHEMA)
    con.execute("INSERT INTO projects (id, code, name, client) VALUES ('P1','PRJ-1','Test','Acme')")
    con.execute(
        "INSERT INTO documents (id, project_id, source, document_type, file_paths, page_count, status)"
        " VALUES ('D1','P1','UPLOAD','INVOICE','[]',1,'APPROVED')"
    )
    con.execute(
        "INSERT INTO doc_headers (document_id, doc_number, total_value) VALUES ('D1','INV-1',118)"
    )
    con.execute(
        "INSERT INTO doc_lines (document_id, line_no, description_raw, quantity, unit, amount)"
        " VALUES ('D1',1,'Cement bags',10,'BAG',100)"
    )
    con.execute(
        "INSERT INTO doc_lines (document_id, line_no, description_raw, quantity, unit, amount)"
        " VALUES ('D1',2,'Cement bags',5,'BAG',50)"
    )
    con.commit()

    project = con.execute("SELECT * FROM projects WHERE id = 'P1'").fetchone()
    wb = load_workbook(build_workbook(con, project))
    assert wb.sheetnames == ["Summary", "Documents", "Line Items", "Materials Rollup"]
    assert wb["Documents"].max_row == 2, "1 header row + 1 document row"
    assert wb["Line Items"].max_row == 3, "1 header row + 2 line rows"
    rollup_row = list(wb["Materials Rollup"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert rollup_row == ("Cement bags", "BAG", 15, 1), rollup_row  # two lines, same material, one doc

    def summary_value(book):
        return next(
            row[1] for row in book["Summary"].iter_rows(values_only=True)
            if row[0] == "Total Value"
        )

    assert summary_value(wb) == 118, summary_value(wb)

    # A rejected document is absent from every sheet — this file states what
    # the project bought, and it leaves the building.
    con.execute(
        "INSERT INTO documents (id, project_id, source, document_type, file_paths, page_count, status)"
        " VALUES ('D2','P1','UPLOAD','INVOICE','[]',1,'REJECTED')"
    )
    con.execute(
        "INSERT INTO doc_headers (document_id, doc_number, total_value) VALUES ('D2','INV-2',999)"
    )
    con.execute(
        "INSERT INTO doc_lines (document_id, line_no, description_raw, quantity, unit, amount)"
        " VALUES ('D2',1,'Basmati Rice',5,'KG',999)"
    )
    con.commit()

    wb2 = load_workbook(build_workbook(con, project))
    body = lambda sheet: [r for r in wb2[sheet].iter_rows(min_row=2, values_only=True)]

    assert wb2["Documents"].max_row == 2, "rejected document dropped from Documents"
    assert wb2["Line Items"].max_row == 3, "its line dropped from Line Items"
    assert all("Basmati" not in str(cell) for row in body("Line Items") for cell in row), \
        body("Line Items")
    assert [r[0] for r in body("Materials Rollup")] == ["Cement bags"], body("Materials Rollup")
    assert summary_value(wb2) == 118, summary_value(wb2)

    # …but the fact that something was dropped is stated, not hidden.
    summary = {r[0]: r[1] for r in wb2["Summary"].iter_rows(values_only=True)}
    assert summary["Rejected (excluded)"] == 1, summary
    assert summary["Total Documents"] == 1, summary
    assert "— REJECTED" not in summary, "the excluded document must not be broken out by status"

    # empty project — no documents at all — must not blow up on empty IN (...)
    con.execute("INSERT INTO projects (id, code, name) VALUES ('P2','PRJ-2','Empty')")
    empty_project = con.execute("SELECT * FROM projects WHERE id = 'P2'").fetchone()
    empty_wb = load_workbook(build_workbook(con, empty_project))
    assert empty_wb["Documents"].max_row == 1, "header row only"

    print("self-check passed")
